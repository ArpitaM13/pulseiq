from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "raw" / "03_PulseIQ_Analytics_Star_Schema_Dataset.xlsx"


def test_source_workbook_has_all_expected_tabs_and_catalog_sizes():
    assert WORKBOOK.exists(), f"Source workbook missing: {WORKBOOK}"
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        assert len(book.sheetnames) == 33
        assert book["metric_dictionary"].max_row - 1 == 28
        assert book["dq_checks_catalog"].max_row - 1 == 14
        assert book["data_contracts"].max_row - 1 == 42
        assert book["dq_known_issues"].max_row - 1 == 11
    finally:
        book.close()


def test_metric_layer_keeps_m28_visible_until_actuals_arrive():
    source = (ROOT / "src" / "metric_layer.py").read_text(encoding="utf-8")
    for metric_id in (f"M{i:02d}" for i in range(1, 29)):
        assert f"'{metric_id}'" in source
    assert "FROM f LEFT JOIN h USING(period, department_id)" in source


def test_warehouse_primary_key_candidates_are_unique_in_source():
    """Every enforced primary key should be valid against the provided workbook."""
    import ast
    import pandas as pd

    source = (ROOT / "src" / "warehouse.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    key_map = None
    for node in tree.body:
        if isinstance(node, ast.AnnAssign) and getattr(node.target, "id", None) == "PRIMARY_KEYS":
            key_map = ast.literal_eval(node.value)
            break
    assert key_map, "PRIMARY_KEYS mapping not found in src/warehouse.py"

    workbook = pd.ExcelFile(WORKBOOK)
    try:
        for sheet, keys in key_map.items():
            frame = pd.read_excel(workbook, sheet_name=sheet)
            assert set(keys).issubset(frame.columns), f"Missing key column on {sheet}: {keys}"
            assert not frame[list(keys)].isna().any(axis=1).any(), f"Null primary-key values in {sheet}: {keys}"
            assert not frame.duplicated(list(keys)).any(), f"Duplicate primary-key values in {sheet}: {keys}"
    finally:
        workbook.close()


def test_task_1_2_4_5_documentation_and_workforce_requirements_exist():
    contracts = (ROOT / "docs" / "data_contract_notes.md").read_text(encoding="utf-8")
    warehouse = (ROOT / "docs" / "warehouse_keys_and_joins.md").read_text(encoding="utf-8")
    phase = (ROOT / "docs" / "phase1_tasks_1_2_4_5.md").read_text(encoding="utf-8")
    workforce = (ROOT / "pages" / "1_Workforce_Headcount.py").read_text(encoding="utf-8")
    assert "fact_task" in contracts and "Nullable" in contracts and "allowed values" in contracts.lower()
    assert "employee_key" in warehouse and "orphan" in warehouse.lower()
    assert all(f"Task {n}" in phase for n in (1, 2, 4, 5))
    assert "date_input" in workforce and "multiselect" in workforce
    assert "Work mode" in workforce and "city_tier" in workforce


def test_metric_layer_avoids_previous_duckdb_reserved_aliases():
    source = (ROOT / "src" / "metric_layer.py").read_text(encoding="utf-8")
    assert "CAST(account_code AS VARCHAR) account_label" in source
    assert "CAST(date_trunc('month', period_end_date) AS DATE) AS period_month" in source
    assert "WITH ar AS (SELECT period month" not in source
    assert " at AS (" not in source
    assert "LEFT JOIN at ON" not in source


def test_streamlit_ui_uses_current_width_api_and_no_department_duplicate_rename():
    root_app = (ROOT / "app.py").read_text(encoding="utf-8")
    workforce = (ROOT / "pages" / "1_Workforce_Headcount.py").read_text(encoding="utf-8")
    helpers = (ROOT / "src" / "dashboard_utils.py").read_text(encoding="utf-8")
    all_ui = root_app + workforce + helpers + "\n".join(p.read_text(encoding="utf-8") for p in (ROOT / "pages").glob("*.py"))
    assert "use_container_width" not in all_ui
    assert 'width="stretch"' in all_ui
    assert 'mix.rename(columns={"department_id":"dimension_label"})' not in root_app


def test_metric_windows_and_monthly_kpi_targets_follow_metric_grain():
    source = (ROOT / "src" / "metric_layer.py").read_text(encoding="utf-8")
    assert "months_in_window = 12" in source
    assert "CAST(date_trunc('month', m.period) AS DATE) AS period_month" in source
    notes = (ROOT / "docs" / "metric_validation.md").read_text(encoding="utf-8").lower()
    assert "full 12-month window" in notes
    assert "monthly targets" in notes


def test_contract_validation_is_explicit_about_partial_source_coverage():
    source = (ROOT / "src" / "contracts.py").read_text(encoding="utf-8")
    docs = (ROOT / "docs" / "data_contract_notes.md").read_text(encoding="utf-8").lower()
    assert "partial_coverage" in source
    assert "cannot be cross-checked" in source
    assert "partially validated" in docs or "partial coverage" in docs


def test_three_core_tables_have_task1_contract_notes():
    notes = (ROOT / "docs" / "core_table_contract_notes.md").read_text(encoding="utf-8")
    for table in ("dim_employee", "fact_payroll", "fact_attendance_monthly"):
        assert f"## `{table}`" in notes
    assert "Data dictionary type" in notes
    assert "Observed blank allowed?" in notes
    assert "Observed nulls" in notes


def test_reports_distinguish_unexpected_columns_from_partial_contract_coverage():
    source = (ROOT / "src" / "reporting.py").read_text(encoding="utf-8")
    assert 'startswith("Unexpected column preserved")' in source
    assert 'partial_contract_coverage_warnings' in source
    assert "core_table_contract_notes.md" in source
